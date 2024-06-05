import sqlite3
import pandas as pd
import os
from docx import Document
from docx.shared import Pt, Cm, Mm
import matplotlib.pyplot as plt


database_path = "db_personas.db"

conn = sqlite3.connect(database_path)
query = """
SELECT p.nombre_completo AS nombre, p.nacionalidad, s.Rol AS rol, s.Sueldo AS salario
FROM personas p
INNER JOIN Salarios s ON p.id_rol = s.id_salarios
"""
df = pd.read_sql_query(query, conn)
conn.close()


def get_person_data(df, full_name):
    person_data = df[df['nombre'] == full_name]
    if person_data.empty:
        raise ValueError(f"No se encontró a la persona con nombre: {full_name}")
    return person_data.iloc[0]


def generar_contrato(date, rol, address, rut, full_name, nationality, birth_date, profession, salary):
    document = Document()
    font = document.styles['Normal'].font
    font.name = 'Book Antiqua'
    sections = document.sections
    for section in sections:
        section.page_height = Cm(35.56)
        section.page_width = Cm(21.59)
        section.top_margin = Cm(0.5)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)
        section.header_distance = Mm(12.7)
        section.footer_distance = Mm(12.7)
    
    header = document.sections[0].header
    paragraph = header.paragraphs[0]
    run = paragraph.add_run()
    run.add_picture("header.png")
    document.add_picture('logo.png')
    
    h = document.add_paragraph('')
    h.add_run('CONTRATO DE PRESTACIÓN DE SERVICIOS A HONORARIOS\n').bold = True
    h.alignment = 1
    font.size = Pt(10)
    h.add_run(
        '............................................................................................................................................................................................................................. ').bold = True
    
    p = document.add_paragraph(f'En Temuco, a {str(date)}, entre la Corporación de Innovación y Desarrollo Tecnológico, Rut 78.898.766-4, representada por su Director General don(a) Roberto Gomez Bolainas, Cédula de Identidad Nº 10.678.990-2, ambos domiciliados en Caupolican 455 de esta ciudad, en adelante la “Corporación” y  {full_name}, de nacionalidad {nationality}, de profesión {profession}, nacido el {birth_date}, con domicilio en {address}, Cédula de Identidad N° {rut}, en adelante, el “Prestador de Servicios”, se ha convenido el siguiente contrato de prestación de servicios a honorarios: \n')
    font.size = Pt(8)
    
    p1 = document.add_paragraph('')
    p1.add_run('PRIMERO        :').bold = True
    p1.add_run('En el marco del acuerdo de servicios profesionales fechado el 11 de noviembre de 2020, establecido entre la Agencia Nacional de Estándares Educativos y la Corporación de Innovación y Desarrollo Tecnológico , y ratificado según la Resolución Exenta N°603 del 23 de noviembre de 2020, la Corporación encarga los servicios profesionales del Prestador de Servicios, para que ejecute la siguiente tarea en el proyecto "Evaluación de competencias específicas y metodologías de aprendizaje artificial 2020, ID 67703-20-JJ90."')
    p1.add_run('\n SEGUNDO        :').bold = True
    p1.add_run('El rol a desempeñar es de '+rol+'.')
    p1.add_run('\n TERCERO        :').bold = True
    p1.add_run('El plazo para la realización de la prestación de servicios encomendada será el '+str(date))
    p1.add_run('\n CUARTO        :').bold = True
    p1.add_run('Por el servicio profesional efectivamente realizado, se pagara un monto bruto variable, el cual corresponderá a cada rol dentro de la empresa capacitación, de acuerdo al siguiente detalle: ')
    
    table = document.add_table(rows=2, cols=2)
    table.alignment = 1
    hdr_cells0 = table.rows[0].cells
    hdr_cells0[0].text = 'Rol'
    hdr_cells0[1].text = 'Monto Bruto'
    hdr_cells = table.rows[1].cells
    hdr_cells[0].text = rol
    hdr_cells[1].text = salary
    
    p1.add_run('\n QUINTO        :').bold = True
    p1.add_run('El Prestador de Servicios acepta el encargo y las condiciones precedentes.')
    p1.add_run('\n SEXTO        :').bold = True
    p1.add_run('El Prestador de Servicios está obligado a mantener la confidencialidad de todos los materiales utilizados, conforme al Acuerdo de Confidencialidad previamente establecido.')
    p1.add_run('\n En comprobante, previa lectura y ratificación, las partes firman.  ').bold = True
    
    table = document.add_table(rows=2, cols=2)
    table.alignment = 1
    hdr_cells0 = table.rows[0].cells[1].add_paragraph()
    r = hdr_cells0.add_run()
    r.add_picture('imagenes/firma.png')
    hdr_cells = table.rows[1].cells
    hdr_cells[0].text = '-----------------------------------------------------------\nEL PRESTADOR DE SERVICIOS'
    hdr_cells[1].text = '-----------------------------------------------------------\np. LA CORPORACION'
    
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    run = paragraph.add_run('Caupolican 0455, Temuco, Chile, www.corpoindet.cl')
    run.add_picture("imagenes/footer1.png")
    
    document.save(f'{full_name}.docx')


def mostrar_menu(df):
    print("Lista de personas disponibles:")
    for index, row in df.iterrows():
        print(f"{index}: {row['nombre']} - {row['rol']} - {row['nacionalidad']} - {row['salario']}")

    seleccion = input("Ingrese los índices de las personas (separados por comas) para generar contratos: ")
    indices = [int(i) for i in seleccion.split(',')]

    for i in indices:
        person_data = df.iloc[i]
        generar_contrato(
            date='2024-06-04',  
            rol=person_data['rol'],
            address='Direccion de ejemplo',  
            rut='12345678-9',  
            full_name=person_data['nombre'],
            nationality=person_data['nacionalidad'],
            birth_date='01-01-1990',  
            profession='Profesion de ejemplo',  
            salary=str(person_data['salario'])
        )
        print(f"Contrato generado para {person_data['nombre']}")


mostrar_menu(df)


promedio_sueldo = df.groupby('rol')['salario'].mean()
promedio_sueldo.plot(kind='bar')
plt.title('Promedio de Sueldo por Profesión')
plt.xlabel('Profesión')
plt.ylabel('Salario Promedio (CLP)')
plt.show()

distribucion_profesiones = df['rol'].value_counts()
distribucion_profesiones.plot(kind='pie', autopct='%1.1f%%')
plt.title('Distribución de Profesiones')
plt.ylabel('')
plt.show()

conteo_nacionalidades = df['nacionalidad'].value_counts()
conteo_nacionalidades.plot(kind='bar')
plt.title('Conteo de Profesionales por Nacionalidad')
plt.xlabel('Nacionalidad')
plt.ylabel('Cantidad')
plt.show()
